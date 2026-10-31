#!/usr/bin/env python3
"""Synthetic, local-only Reader V1 acceptance checks."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
PDFJS_MARKER = 'const PDFJS_URL = "/vendor/reader/pdfjs/pdf.mjs"'


def tiny_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    body = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, value in enumerate(objects, start=1):
        offsets.append(len(body))
        body.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        body.extend(f"{offset:010d} 00000 n \n".encode())
    body.extend(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(body)


def write_docx(path: Path) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/xml"/><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Synthetic Reader DOCX</w:t></w:r></w:p></w:body></w:document>')


def main() -> int:
    import reader_extractors as reader

    package = json.loads((ROOT / "package.json").read_text(encoding="utf-8"))
    reader_js = (ROOT / "static" / "reader.js").read_text(encoding="utf-8")
    app_js = (ROOT / "static" / "app.js").read_text(encoding="utf-8")
    extractor_source = (ROOT / "reader_extractors.py").read_text(encoding="utf-8")
    config = json.loads((ROOT / "src-tauri" / "tauri.conf.json").read_text(encoding="utf-8"))
    checks = {
        "pdfjs_exact_pin": package["dependencies"].get("pdfjs-dist") == "6.4.299",
        "prism_exact_pin": package["dependencies"].get("prismjs") == "1.30.0",
        "mac_ocr_exact_pin": package["optionalDependencies"].get("mac-ocr") == "1.1.1",
        "local_pdfjs_assets": PDFJS_MARKER in reader_js and "https://" not in reader_js,
        "pdfjs_security_flags": "isEvalSupported: false" in reader_js and "enableXfa: false" in reader_js,
        "bounded_page_rendering": "MAX_RENDERED_PAGES = 7" in reader_js and "trimRendered" in reader_js,
        "obsolete_render_cancellation": "RenderingCancelledException" in reader_js and "task.cancel()" in reader_js,
        "fit_width_and_page": "fitWidth()" in reader_js and "fitPage()" in reader_js,
        "reader_route_restoration": "#/file/" in app_js and "onPageChange" in reader_js,
        "no_iframe_reader": "<iframe" not in reader_js and "StudyHubReader.open" in app_js,
        "no_untrusted_html_rendering": "code.textContent =" in reader_js,
        "ocr_packaged_resource": "reader-tools/mac-ocr" in json.dumps(config),
        "markitdown_helper_packaged_resource": "reader-tools/studyhub-markitdown" in json.dumps(config),
        "markitdown_plugins_disabled": "enable_plugins=False" in extractor_source,
        "helper_path_not_in_argv": "[str(helper)]" in extractor_source and 'stdin=subprocess.PIPE' in extractor_source,
    }
    with tempfile.TemporaryDirectory(prefix="studyhub-reader-v1-") as raw:
        root = Path(raw)
        pdf = root / "reader.pdf"
        pdf.write_bytes(tiny_pdf("Synthetic Reader PDF"))
        before_pdf = hashlib.sha256(pdf.read_bytes()).hexdigest()
        node = subprocess.run(
            ["node", "--input-type=module", "-e", "import fs from 'node:fs'; import {getDocument} from './node_modules/pdfjs-dist/legacy/build/pdf.mjs'; const task=getDocument({data:new Uint8Array(fs.readFileSync(process.argv[1])),isEvalSupported:false}); const d=await task.promise; const p=await d.getPage(1); const t=await p.getTextContent(); if(d.numPages!==1||!t.items.map(x=>x.str).join(' ').includes('Synthetic Reader PDF'))process.exit(2); await task.destroy();", str(pdf)],
            cwd=ROOT, capture_output=True, text=True,
        )
        checks["real_synthetic_pdf_pdfjs"] = node.returncode == 0
        checks["pdf_source_unchanged"] = before_pdf == hashlib.sha256(pdf.read_bytes()).hexdigest()

        docx = root / "reader.docx"
        write_docx(docx)
        pptx = root / "reader.pptx"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = "Synthetic Reader PPTX"
        presentation.save(pptx)
        xlsx = root / "reader.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "Synthetic Reader XLSX"
        workbook.save(xlsx)
        for path, expected in ((docx, "Synthetic Reader DOCX"), (pptx, "Synthetic Reader PPTX"), (xlsx, "Synthetic Reader XLSX")):
            before = hashlib.sha256(path.read_bytes()).hexdigest()
            result = reader.extract_with_markitdown(path, [root])
            checks[f"markitdown_{path.suffix[1:]}"] = result.method == "markitdown" and expected in result.text
            checks[f"unchanged_{path.suffix[1:]}"] = before == hashlib.sha256(path.read_bytes()).hexdigest()

        helper = root / "studyhub-markitdown"
        helper.write_text(
            "#!/bin/sh\n"
            "[ \"$#\" -eq 0 ] || exit 7\n"
            "IFS= read -r request\n"
            "printf '%s' \"$request\" | grep -q '\"operation\": \"convert\"' || exit 8\n"
            "printf '%s\\n' '{\"ok\":true,\"text\":\"Synthetic helper text\"}'\n",
            encoding="utf-8",
        )
        helper.chmod(0o755)
        previous_helper = os.environ.get(reader.MARKITDOWN_HELPER_ENV)
        os.environ[reader.MARKITDOWN_HELPER_ENV] = str(helper)
        try:
            helper_result = reader.extract_with_markitdown(docx, [root])
        finally:
            if previous_helper is None:
                os.environ.pop(reader.MARKITDOWN_HELPER_ENV, None)
            else:
                os.environ[reader.MARKITDOWN_HELPER_ENV] = previous_helper
        checks["packaged_helper_stdin_transport"] = helper_result.text == "Synthetic helper text"
        reader._stop_markitdown_helper()

        image = root / "synthetic.png"
        image.write_bytes(b"synthetic-image-input")
        fake = root / "mac-ocr"
        fake.write_text('#!/bin/sh\nprintf \'%s\\n\' \'{"page":1,"pageCount":1,"text":"Synthetic OCR text"}\'\n', encoding="utf-8")
        fake.chmod(0o755)
        before = hashlib.sha256(image.read_bytes()).hexdigest()
        ocr = reader.extract_with_mac_ocr(image, [root], fake, timeout_seconds=2)
        checks["ocr_bounded_adapter"] = ocr.ocr_used and ocr.page_labels == ("Page 1",) and "Synthetic OCR text" in ocr.text
        checks["ocr_source_unchanged"] = before == hashlib.sha256(image.read_bytes()).hexdigest()
        try:
            reader.extract_with_mac_ocr(Path("https://example.invalid/input.png"), [root], fake)
            checks["ocr_url_rejected"] = False
        except (ValueError, OSError):
            checks["ocr_url_rejected"] = True

    for name, passed in checks.items():
        print(f"{name}: {'PASS' if passed else 'FAIL'}")
    return 0 if all(checks.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
